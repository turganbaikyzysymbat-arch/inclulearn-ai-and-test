import io
import json
import hashlib
import hmac
import secrets
import os
import re
import sqlite3
import urllib.request
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path

from flask import Flask, jsonify, render_template, request, session, redirect, url_for


BASE_DIR = Path(__file__).resolve().parent
DB_DIR = BASE_DIR / "instance"
DB_PATH = DB_DIR / "learn4all.db"
DB_DIR.mkdir(exist_ok=True)

app = Flask(__name__, template_folder="templates", static_folder="static")
app.secret_key = os.getenv("FLASK_SECRET_KEY") or secrets.token_hex(32)
app.config.update(SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax",
                  SESSION_COOKIE_SECURE=os.getenv("RENDER", "").lower() == "true")
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

STOP_WORDS = set(
    "және мен бұл үшін туралы немесе бір оның болып арқылы да де әрі қалай деген "
    "the and of to a in is it that for with as on are was this be by an or from at "
    "и в на это как для что из с по".split()
)

LANGUAGE_NAMES = {
    "kk": "қазақ тілі",
    "ru": "русский язык",
    "en": "English",
}


@contextmanager
def db():
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        yield connection
        connection.commit()
    finally:
        connection.close()


def init_db():
    DB_DIR.mkdir(exist_ok=True)
    with db() as connection:
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                email TEXT NOT NULL UNIQUE COLLATE NOCASE,
                password_hash TEXT,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS task_sets (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                title TEXT NOT NULL,
                tasks_json TEXT NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS materials (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                source_text TEXT NOT NULL,
                file_name TEXT,
                language TEXT DEFAULT 'kk',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS adaptations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                material_id INTEGER NOT NULL UNIQUE,
                result_json TEXT NOT NULL,
                engine TEXT NOT NULL DEFAULT 'local',
                created_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS questions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                material_id INTEGER NOT NULL,
                question TEXT NOT NULL,
                answer TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS progress (
                material_id INTEGER PRIMARY KEY,
                completed INTEGER NOT NULL DEFAULT 0,
                score INTEGER NOT NULL DEFAULT 0,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(material_id) REFERENCES materials(id) ON DELETE CASCADE
            );
            """
        )


def now():
    return datetime.now().isoformat(timespec="seconds")


def clean_text(value):
    value = str(value or "").replace("\x00", " ").replace("\r", "")
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def split_sentences(text):
    return [
        item.strip()
        for item in re.split(
            r"(?<=[.!?])\s+(?=[А-ЯӘҒҚҢӨҰҮҺІA-Z0-9])", clean_text(text)
        )
        if item.strip()
    ]


def get_words(text):
    return re.findall(
        r"[а-яәіңғқөұүһіa-z0-9'-]+", clean_text(text).lower(), flags=re.I
    )


def keywords(text, limit=8):
    counts = {}
    for word in get_words(text):
        if len(word) >= 4 and word not in STOP_WORDS:
            counts[word] = counts.get(word, 0) + 1
    return [
        word
        for word, _ in sorted(counts.items(), key=lambda item: (-item[1], item[0]))[
            :limit
        ]
    ]


def select_summary(text, limit=4):
    sentence_list = split_sentences(text)
    if len(sentence_list) <= limit:
        return sentence_list
    terms = keywords(text, limit)
    scored = []
    for position, sentence in enumerate(sentence_list):
        score = sum(2 for term in terms if term in sentence.lower())
        score += 1 if position == 0 else 0
        scored.append((score, position, sentence))
    return [
        row[2]
        for row in sorted(sorted(scored, reverse=True)[:limit], key=lambda row: row[1])
    ]


def local_adaptation(text, language="kk"):
    source = clean_text(text)
    terms = keywords(source)
    summary_sentences = select_summary(source)
    topic = ", ".join(terms[:3]) or "негізгі тақырып"
    short = "\n".join(
        f"• {sentence[:180]}{'…' if len(sentence) > 180 else ''}"
        for sentence in summary_sentences
    )
    steps = "\n".join(
        f"{index}. {sentence}"
        for index, sentence in enumerate(split_sentences(source)[:8], 1)
    ) or "1. Мәтінді оқыңыз.\n2. Негізгі ұғымдарды белгілеңіз.\n3. Мысал арқылы тексеріңіз."
    examples = "\n".join(
        f"{index}. «{term}» ұғымын күнделікті өмірдегі жағдаймен байланыстырып, өз мысалыңызды жазыңыз."
        for index, term in enumerate(terms[:4], 1)
    ) or "1. Тақырыпқа қатысты қарапайым жағдай ойластырыңыз.\n2. Оны негізгі ұғыммен байланыстырыңыз."
    tasks = (
        f"1. «{terms[0] if terms else 'негізгі ұғым'}» ұғымын өз сөзіңізбен түсіндіріңіз.\n"
        f"2. {terms[1] if len(terms) > 1 else 'тақырып'} бойынша бір мысал келтіріңіз.\n"
        "3. Мәтіндегі екі негізгі ойдың байланысын жазыңыз.\n"
        "4. Жауабыңыздан негізгі терминдерді белгілеңіз."
    )
    plain = re.sub(r"\([^)]*\)", "", source).replace(";", ".")
    plain = "\n\n".join(
        sentence[:200] + ("…" if len(sentence) > 200 else "")
        for sentence in split_sentences(plain)
    )
    nodes = (terms or ["Негізгі ұғым", "Ереже", "Мысал", "Нәтиже"])[:5]
    return {
        "simplified": short or source[:600],
        "steps": steps,
        "examples": examples,
        "diagram": {
            "title": topic,
            "nodes": nodes,
            "edges": [[nodes[index], nodes[index + 1]] for index in range(len(nodes) - 1)],
        },
        "large": plain or source,
        "summary": " ".join(summary_sentences) or source[:500],
        "tasks": tasks,
        "keywords": terms,
        "explanation": (
            f"Бұл материалдың негізгі бағыты: {topic}. Алдымен анықтаманы түсініңіз, "
            "кейін оның қолданылуын мысалмен тексеріңіз."
        ),
        "language": language,
    }


def ai_configured():
    return bool(os.getenv("OPENAI_API_KEY") or os.getenv("AI_INTEGRATIONS_OPENAI_API_KEY"))


def ask_model(system, prompt):
    api_key = os.getenv("OPENAI_API_KEY") or os.getenv("AI_INTEGRATIONS_OPENAI_API_KEY")
    if not api_key:
        return None
    endpoint = (
        os.getenv("OPENAI_BASE_URL", "https://api.openai.com/v1").rstrip("/")
        + "/chat/completions"
    )
    payload = json.dumps(
        {
            "model": os.getenv("OPENAI_MODEL", "gpt-4o-mini"),
            "temperature": 0.2,
            "messages": [
                {"role": "system", "content": system},
                {"role": "user", "content": prompt},
            ],
        }
    ).encode()
    model_request = urllib.request.Request(
        endpoint,
        data=payload,
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"},
        method="POST",
    )
    with urllib.request.urlopen(model_request, timeout=45) as response:
        data = json.loads(response.read().decode())
    return data["choices"][0]["message"]["content"]


def extract_file(file):
    extension = Path(file.filename or "").suffix.lower()
    allowed = {".txt", ".md", ".csv", ".html", ".pdf", ".docx", ".pptx"}
    if extension not in allowed:
        raise ValueError("PDF, DOCX, PPTX, TXT, MD, CSV немесе HTML файлдарын қолданыңыз.")
    data = file.read()
    if extension in {".txt", ".md", ".csv", ".html"}:
        return clean_text(data.decode("utf-8", errors="ignore"))
    if extension == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return clean_text("\n\n".join(page.extract_text() or "" for page in reader.pages))
    if extension == ".docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        return clean_text("\n\n".join(paragraph.text for paragraph in document.paragraphs))
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    slides = []
    for slide in presentation.slides:
        slide_text = [
            shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()
        ]
        if slide_text:
            slides.append(" ".join(slide_text))
    return clean_text("\n\n".join(slides))


def local_answer(question, source):
    query = keywords(question, 8)
    matched = []
    for position, sentence in enumerate(split_sentences(source)):
        score = sum(1 for term in query if term in sentence.lower())
        if score:
            matched.append((score, -position, sentence))
    matched = [row[2] for row in sorted(matched, reverse=True)[:3]]
    if not matched:
        return "Бұл сұраққа жүктелген материалдан нақты жауап табылмады. Сұрақты мәтіндегі негізгі терминдермен нақтылап көріңіз."
    return "Материалға сүйенген жауап:\n\n" + "\n".join(f"• {sentence}" for sentence in matched)


def material_json(row, adaptation=None, progress=None, preview=False):
    result = dict(row)
    if preview:
        result["source_text"] = result["source_text"][:180]
    result["adaptation"] = json.loads(adaptation["result_json"]) if adaptation else None
    result["progress"] = dict(progress) if progress else {
        "completed": 0,
        "score": 0,
    }
    return result


@app.before_request
def require_teacher_login():
    allowed = {"index", "register", "login", "logout", "health", "static"}
    if request.endpoint in allowed or request.endpoint is None:
        return None
    if request.path.startswith("/api/") and request.path != "/api/health" and not session.get("user_id"):
        return jsonify({"error": "Алдымен мұғалім аккаунтына кіріңіз."}), 401
    return None


@app.get("/")
def index():
    user = None
    if session.get("user_id"):
        user = {"id": session["user_id"], "name": session.get("user_name", "Мұғалім"),
                "email": session.get("user_email", "")}
    return render_template("index.html", current_user=user)


def password_hash(password, salt=None):
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 240000).hex()
    return f"{salt}${digest}"


def password_matches(password, stored):
    if not stored or "$" not in stored:
        return False
    salt, expected = stored.split("$", 1)
    actual = password_hash(password, salt).split("$", 1)[1]
    return hmac.compare_digest(actual, expected)


def sign_in(user):
    session.clear()
    session["user_id"] = user["id"]
    session["user_name"] = user["name"]
    session["user_email"] = user["email"]


@app.post("/auth/register")
def register():
    body = request.get_json(silent=True) or {}
    name = clean_text(body.get("name"))
    email = clean_text(body.get("email")).lower()
    password = str(body.get("password") or "")
    if len(name) < 2 or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        return jsonify({"error": "Аты-жөніңізді және дұрыс электрондық поштаны енгізіңіз."}), 400
    if len(password) < 8:
        return jsonify({"error": "Құпиясөз кемінде 8 таңбадан тұруы керек."}), 400
    try:
        with db() as connection:
            cursor = connection.execute("INSERT INTO users(name,email,password_hash,created_at) VALUES(?,?,?,?)",
                                        (name, email, password_hash(password), now()))
            user = {"id": cursor.lastrowid, "name": name, "email": email}
        sign_in(user)
        return jsonify({"ok": True, "user": user})
    except sqlite3.IntegrityError:
        return jsonify({"error": "Бұл электрондық пошта тіркелген. Кіру бөлімін қолданыңыз."}), 409


@app.post("/auth/login")
def login():
    body = request.get_json(silent=True) or {}
    email = clean_text(body.get("email")).lower()
    password = str(body.get("password") or "")
    with db() as connection:
        user = connection.execute("SELECT * FROM users WHERE email = ?", (email,)).fetchone()
    if not user or not password_matches(password, user["password_hash"]):
        return jsonify({"error": "Пошта немесе құпиясөз дұрыс емес."}), 401
    sign_in(user)
    return jsonify({"ok": True, "user": {"id": user["id"], "name": user["name"], "email": user["email"]}})


@app.post("/auth/logout")
def logout():
    session.clear()
    return jsonify({"ok": True})


@app.post("/api/generate-task-set")
def generate_task_set():
    """Create a teacher task set from pasted text or an uploaded learning file."""
    title = clean_text(request.form.get("title") or "AI құрастырған тапсырмалар")[:160]
    source = clean_text(request.form.get("source") or "")
    uploaded = request.files.get("file")
    if uploaded and uploaded.filename:
        try:
            source = extract_file(uploaded)
        except Exception as exc:
            return jsonify({"error": f"Файлды оқу мүмкін болмады: {exc}"}), 400
    if len(source.strip()) < 80:
        return jsonify({"error": "Кемінде 80 таңбадан тұратын материал енгізіңіз немесе оқылатын файл жүктеңіз."}), 400
    try:
        count = max(5, min(20, int(request.form.get("count", "10"))))
    except ValueError:
        count = 10
    kind = request.form.get("kind", "mixed")
    level = request.form.get("level", "medium")
    types = {"quiz", "multi", "truefalse", "matching", "scenario", "open", "fill", "order"}
    chosen = [kind] if kind in types else ["quiz", "truefalse", "matching", "scenario", "open", "fill", "order"]
    system = "Сен мұғалімдерге арналған бағалау тапсырмаларын құратын қазақтілді педагогикалық көмекшісің. Тек берілген материалға сүйен. Жауап кілттері нақты және материалмен расталсын. Тек жарамды JSON қайтар."
    prompt = (f"Материал негізінде {count} тапсырма құрастыр. Деңгейі: {level}. Рұқсат етілген түрлер: {chosen}. JSON форматы: tasks массиві; әр элементте type, question және сол түрге сай options, answers, answer, pairs, expected немесе items болсын. quiz: 4 нұсқа және бір дұрыс жауап; multi: 4 нұсқа және кемінде 2 дұрыс жауап; truefalse: answer true/false; matching: кемінде 3 жұп; fill: answer; order: items; open/scenario: expected. Дұрыс жауапты жорамалмен ойдан қоспа.\nМАТЕРИАЛ:\n{source[:18000]}")
    tasks = None
    ai_generated = False
    if ai_configured():
        try:
            raw = ask_model(system, prompt)
            match = re.search(r"\{.*\}", raw or "", re.S)
            if match:
                tasks = json.loads(match.group(0)).get("tasks")
                ai_generated = isinstance(tasks, list) and bool(tasks)
        except Exception:
            tasks = None
    # Source-grounded offline fallback. It never pretends to be a language model.
    if not isinstance(tasks, list) or not tasks:
        sentences = [x.strip() for x in split_sentences(source) if len(x.strip()) >= 20]
        if not sentences:
            sentences = [source.strip()[:500]]
        terms = keywords(source, max(12, count * 2))
        tasks = []
        for i in range(count):
            sentence = sentences[i % len(sentences)]
            term = terms[i % len(terms)] if terms else sentence.split()[0].strip(".,:;!?«»\"'()")
            task_type = kind if kind in types else ["quiz", "truefalse", "matching", "scenario", "open", "fill", "order"][i % 7]
            if task_type == "quiz":
                # The exact source sentence is the only answer; distractors are visibly altered excerpts.
                correct = sentence[:220]
                wrongs = []
                # Alter the source sentence so distractors are not other true source sentences.
                tokens = correct.split()
                for j in range(1, 4):
                    if len(tokens) > 2:
                        altered = tokens[:]
                        idx = min(len(altered)-1, max(1, (i + j) % len(altered)))
                        replacement = (terms[(i + j) % len(terms)] if terms else "басқа")
                        if replacement.casefold() == altered[idx].strip(".,:;!?«»\"'").casefold():
                            replacement = "өзгертілген"
                        altered[idx] = replacement
                        candidate = " ".join(altered)[:220]
                    else:
                        candidate = correct + f" (өзгертілген нұсқа {j})"
                    if candidate not in wrongs and candidate != correct:
                        wrongs.append(candidate)
                while len(wrongs) < 3:
                    wrongs.append(correct + f" (өзгертілген нұсқа {len(wrongs)+1})")
                options = [correct] + wrongs[:3]
                # Deterministic rotation avoids the correct option always appearing first.
                shift = i % 4
                options = options[shift:] + options[:shift]
                tasks.append({"type":"quiz","question":"Қай нұсқа материалдағы сөйлемді дәл береді?","options":options,"answers":[correct]})
            elif task_type == "multi":
                present = list(dict.fromkeys(terms[:max(2, min(3, len(terms)))]))
                if len(present) < 2:
                    present = [term, sentence.split()[0]]
                distractors = [f"{term}_жоқ_{j}" for j in range(1, 5)]
                options = (present + distractors)[:4]
                if len(options) < 4: options += [f"бөгде_ұғым_{j}" for j in range(4-len(options))]
                tasks.append({"type":"multi","question":"Материалда кездесетін негізгі терминдерді таңдаңыз.","options":options,"answers":present[:min(3,len(options))]})
            elif task_type == "truefalse":
                tasks.append({"type":"truefalse","question":f"Материалда мына ой берілген: «{sentence[:220]}».","answer":"true"})
            elif task_type == "matching":
                pairs=[]
                for j, term2 in enumerate(terms[:4]):
                    found=next((sent for sent in sentences if term2.lower() in sent.lower()), sentence)
                    pair={"left":term2,"right":found[:220]}
                    if pair not in pairs: pairs.append(pair)
                if len(pairs)<2:
                    pairs=[{"left":f"Үзінді {j+1}","right":sentences[j % len(sentences)][:220]} for j in range(2)]
                tasks.append({"type":"matching","question":"Терминдерді олар кездесетін мәтін үзінділерімен сәйкестендіріңіз.","pairs":pairs[:4]})
            elif task_type == "scenario":
                tasks.append({"type":"scenario","question":f"Жағдаят: «{sentence[:180]}» деген ақпаратты сабақтағы нақты жағдайда қалай қолдануға болады? Материалға сүйеніп түсіндіріңіз.","expected":"Шешім материалдағы ойға сүйенеді; жауапта әрекет пен оның себебі түсіндіріледі."})
            elif task_type == "fill":
                masked = re.sub(re.escape(term), "________", sentence, count=1, flags=re.I)
                if masked == sentence:
                    term = sentence.split()[0].strip(".,:;!?«»\"'()") or term
                    masked = sentence.replace(term, "________", 1)
                tasks.append({"type":"fill","question":f"Бос орынды материалдағы сөзбен толтырыңыз: «{masked[:260]}»","answer":term})
            elif task_type == "order":
                items = sentences[:min(4, len(sentences))]
                if len(items)<2: items = [sentence, f"Мәтіндегі келесі ой: {sentence}"]
                # Store intended order for checking; shuffle display order client-side is not required.
                tasks.append({"type":"order","question":"Мәтіндегі ойлардың логикалық ретін анықтаңыз.","items":items})
            else:
                tasks.append({"type":"open","question":f"Материалға сүйеніп түсіндіріңіз: «{sentence[:220]}» Негізгі ойды өз сөзіңізбен баяндаңыз.","expected":"Жауап материалдағы негізгі ойды бұрмаламай түсіндіреді және мәтіндегі ақпаратпен байланыстырады."})
    cleaned=[]
    for task in tasks[:count]:
        if not isinstance(task, dict) or not clean_text(task.get("question")): continue
        typ=task.get("type")
        if typ not in types: typ="open"
        if kind in types and typ != kind: continue
        item={"type":typ,"question":clean_text(task.get("question"))}
        if typ in {"quiz","multi"}:
            item["options"]=[clean_text(x) for x in task.get("options",[]) if clean_text(x)]
            item["answers"]=[clean_text(x) for x in task.get("answers",[]) if clean_text(x)]
            if len(item["options"])<2 or not item["answers"] or any(x not in item["options"] for x in item["answers"]): continue
            if typ=="quiz" and len(item["answers"])!=1: continue
        elif typ=="truefalse":
            if task.get("answer") not in (True,False,"true","false"): continue
            item["answer"]="true" if task.get("answer") in (True,"true") else "false"
        elif typ=="matching":
            item["pairs"]=[{"left":clean_text(x.get("left")),"right":clean_text(x.get("right"))} for x in task.get("pairs",[]) if isinstance(x,dict) and clean_text(x.get("left")) and clean_text(x.get("right"))]
            if len(item["pairs"])<2: continue
        elif typ in {"open","scenario"}:
            item["expected"]=clean_text(task.get("expected")) or "Жауап материалға сай болуы және негізгі ойды түсіндіруі керек."
        elif typ=="fill":
            item["answer"]=clean_text(task.get("answer"))
            if not item["answer"]: continue
        elif typ=="order":
            item["items"]=[clean_text(x) for x in task.get("items",[]) if clean_text(x)]
            if len(item["items"])<2: continue
        cleaned.append(item)
    if not cleaned:
        return jsonify({"error":"Материалдан сенімді тапсырмалар жасалмады. Материалды нақтырақ енгізіп, қайта көріңіз."}), 422
    if len(cleaned) < count:
        return jsonify({"error":f"Тек {len(cleaned)} жарамды тапсырма жасалды. Жауап кілті толық емес тапсырмалар алынып тасталды; басқа материалмен немесе AI кілтін тексеріп қайта көріңіз."}), 422
    return jsonify({"title":title,"tasks":cleaned,"source_length":len(source),"ai_used":ai_generated,"engine":"ai" if ai_generated else "source-based-local"})


@app.post("/api/task-sets")
def create_task_set():
    body = request.get_json(silent=True) or {}
    title = clean_text(body.get("title"))
    tasks = body.get("tasks")
    if not title or not isinstance(tasks, list) or not tasks:
        return jsonify({"error": "Жұмыс атауын және кемінде бір тапсырманы енгізіңіз."}), 400
    allowed_types = {"quiz", "multi", "truefalse", "matching", "scenario", "open", "fill", "order"}
    for index, task in enumerate(tasks, 1):
        if task.get("type") not in allowed_types or not clean_text(task.get("question")):
            return jsonify({"error": f"{index}-тапсырмада түрі немесе шарты жоқ."}), 400
        if task["type"] in {"quiz", "multi"}:
            options = task.get("options", [])
            answers = task.get("answers", [])
            if len(options) < 2 or not answers or any(a not in options for a in answers):
                return jsonify({"error": f"{index}-тест үшін нұсқалар мен дұрыс жауаптарды тексеріңіз."}), 400
            if task["type"] == "quiz" and len(answers) != 1:
                return jsonify({"error": f"{index}-тапсырмада бір ғана дұрыс жауап болуы керек."}), 400
        if task["type"] == "truefalse" and task.get("answer") not in ["true", "false"]:
            return jsonify({"error": f"{index}-тапсырма үшін дұрыс/бұрыс жауабын таңдаңыз."}), 400
        if task["type"] == "matching" and len(task.get("pairs", [])) < 2:
            return jsonify({"error": f"{index}-сәйкестендіруге кемінде 2 жұп керек."}), 400
        if task["type"] in {"scenario", "open"} and not clean_text(task.get("expected", "")):
            return jsonify({"error": f"{index}-ашық/жағдаяттық тапсырмаға күтілетін жауап немесе критерий жазыңыз."}), 400
        if task["type"] == "fill" and not clean_text(task.get("answer", "")):
            return jsonify({"error": f"{index}-бос орынды толтыруға жауап енгізіңіз."}), 400
        if task["type"] == "order" and len(task.get("items", [])) < 2:
            return jsonify({"error": f"{index}-реттеу тапсырмасына кемінде 2 қадам енгізіңіз."}), 400
    with db() as connection:
        cursor = connection.execute("INSERT INTO task_sets(user_id,title,tasks_json,created_at) VALUES(?,?,?,?)",
            (session.get("user_id"), title, json.dumps(tasks, ensure_ascii=False), now()))
        task_set_id = cursor.lastrowid
    return jsonify({"id": task_set_id, "title": title, "tasks": tasks, "created_at": now()}), 201


@app.get("/api/task-sets")
def list_task_sets():
    with db() as connection:
        rows = connection.execute("SELECT id,title,created_at,tasks_json FROM task_sets WHERE user_id=? ORDER BY id DESC",
                                  (session.get("user_id"),)).fetchall()
    return jsonify({"task_sets": [{"id": r["id"], "title": r["title"], "created_at": r["created_at"],
                                  "count": len(json.loads(r["tasks_json"]))} for r in rows]})


@app.get("/api/task-sets/<int:set_id>")
def get_task_set(set_id):
    with db() as connection:
        row = connection.execute("SELECT * FROM task_sets WHERE id=? AND user_id=?",
                                 (set_id, session.get("user_id"))).fetchone()
    if not row:
        return jsonify({"error": "Жұмыс табылмады."}), 404
    return jsonify({"id": row["id"], "title": row["title"], "tasks": json.loads(row["tasks_json"])})


@app.post("/api/task-sets/<int:set_id>/grade")
def grade_task_set(set_id):
    body = request.get_json(silent=True) or {}
    with db() as connection:
        row = connection.execute("SELECT * FROM task_sets WHERE id=? AND user_id=?",
                                 (set_id, session.get("user_id"),)).fetchone()
    if not row:
        return jsonify({"error": "Жұмыс табылмады."}), 404
    tasks = json.loads(row["tasks_json"])
    answers = body.get("answers", [])
    results, earned, total = [], 0, 0
    for index, task in enumerate(tasks):
        response = answers[index] if index < len(answers) else None
        kind = task["type"]
        correct = None
        if kind in {"quiz", "multi"}:
            given = response if isinstance(response, list) else [str(response or "")]
            correct = set(given) == set(task["answers"])
        elif kind == "truefalse":
            correct = str(response).lower() == task["answer"]
        elif kind == "fill":
            accepted = [x.strip().casefold() for x in task["answer"].split("|")]
            correct = str(response or "").strip().casefold() in accepted
        elif kind == "matching":
            expected = {str(i): str(pair["right"]) for i, pair in enumerate(task["pairs"])}
            correct = isinstance(response, dict) and all(str(response.get(k, "")).strip().casefold() == v.strip().casefold() for k,v in expected.items())
        elif kind == "order":
            correct = isinstance(response, list) and response == task.get("items", [])
        if correct is not None:
            total += 1
            earned += int(correct)
        results.append({"index": index, "type": kind, "correct": correct,
                        "expected": task.get("answers", task.get("answer", task.get("expected", task.get("items", []))))})
    return jsonify({"earned": earned, "total_auto_graded": total, "results": results,
                    "score": round(earned * 100 / total) if total else None,
                    "note": "Ашық және жағдаяттық жауаптар мұғалімнің критерийі бойынша қолмен бағаланады."})



@app.get("/api/health")
def health():
    return jsonify({"ok": True, "database": DB_PATH.exists(), "ai_configured": ai_configured()})


@app.get("/api/materials")
def list_materials():
    query = clean_text(request.args.get("q", ""))
    with db() as connection:
        if query:
            rows = connection.execute(
                "SELECT * FROM materials WHERE title LIKE ? OR source_text LIKE ? "
                "ORDER BY updated_at DESC LIMIT 100",
                (f"%{query}%", f"%{query}%"),
            ).fetchall()
        else:
            rows = connection.execute(
                "SELECT * FROM materials ORDER BY updated_at DESC LIMIT 100"
            ).fetchall()
        materials = []
        for row in rows:
            progress = connection.execute(
                "SELECT completed, score, updated_at FROM progress WHERE material_id = ?",
                (row["id"],),
            ).fetchone()
            materials.append(material_json(row, progress=progress, preview=True))
    return jsonify({"materials": materials})


@app.get("/api/materials/<int:material_id>")
def get_material(material_id):
    with db() as connection:
        row = connection.execute(
            "SELECT * FROM materials WHERE id = ?", (material_id,)
        ).fetchone()
        adaptation = connection.execute(
            "SELECT * FROM adaptations WHERE material_id = ?", (material_id,)
        ).fetchone()
        progress = connection.execute(
            "SELECT completed, score, updated_at FROM progress WHERE material_id = ?",
            (material_id,),
        ).fetchone()
    if not row:
        return jsonify({"error": "Материал табылмады."}), 404
    return jsonify(material_json(row, adaptation, progress))


@app.post("/api/materials")
def create_material():
    try:
        body = request.get_json(silent=True) if request.is_json else {}
        title = clean_text(request.form.get("title") or body.get("title", ""))
        language = request.form.get("language") or body.get("language") or "kk"
        file = request.files.get("file")
        text = request.form.get("text", "") or body.get("text", "")
        file_name = file.filename if file else ""
        if file:
            text = extract_file(file)
        text = clean_text(text)
        if not text:
            return jsonify({"error": "Материал мәтінін енгізіңіз немесе файл жүктеңіз."}), 400
        if not title:
            title = Path(file_name).stem if file_name else "Жаңа оқу материалы"
        timestamp = now()
        with db() as connection:
            cursor = connection.execute(
                "INSERT INTO materials(title, source_text, file_name, language, created_at, updated_at) "
                "VALUES (?, ?, ?, ?, ?, ?)",
                (title, text, file_name, language, timestamp, timestamp),
            )
            material_id = cursor.lastrowid
        result = await_adaptation(material_id, text, language)
        return jsonify(
            {"id": material_id, "title": title, "source_text": text, "adaptation": result}
        ), 201
    except Exception as error:
        return jsonify({"error": str(error)}), 400


def await_adaptation(material_id, text, language):
    result = local_adaptation(text, language)
    engine = "local"
    if ai_configured():
        try:
            prompt = (
                "Берілген материалды инклюзивті оқытуға бейімде. Тек материалдағы ақпаратты пайдалан. "
                f"Нәтижені {LANGUAGE_NAMES.get(language, 'қазақ тілі')} тілінде бер. "
                "Төмендегі JSON кілттерін сақта: simplified, steps, examples, explanation, summary, tasks."
                f"\nМатериал:\n{text}"
            )
            response = ask_model(
                "Сен IncluLearn AI платформасындағы Ayla AI ассистентісің. Қысқа, түсінікті жауап бер.",
                prompt,
            )
            parsed = json.loads(response)
            result.update(
                {key: parsed[key] for key in parsed if key in result and isinstance(parsed[key], str)}
            )
            engine = "ai"
        except Exception:
            engine = "local"
    with db() as connection:
        connection.execute(
            "INSERT INTO adaptations(material_id, result_json, engine, created_at) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(material_id) DO UPDATE SET result_json=excluded.result_json, "
            "engine=excluded.engine, created_at=excluded.created_at",
            (material_id, json.dumps(result, ensure_ascii=False), engine, now()),
        )
        connection.execute(
            "UPDATE materials SET updated_at = ? WHERE id = ?", (now(), material_id)
        )
    return result


@app.put("/api/materials/<int:material_id>")
def update_material(material_id):
    body = request.get_json(silent=True) or {}
    title = clean_text(body.get("title"))
    text = clean_text(body.get("text"))
    if not title or not text:
        return jsonify({"error": "Атауы мен мәтіні бос болмауы керек."}), 400
    with db() as connection:
        exists = connection.execute(
            "SELECT id FROM materials WHERE id = ?", (material_id,)
        ).fetchone()
        if not exists:
            return jsonify({"error": "Материал табылмады."}), 404
        connection.execute(
            "UPDATE materials SET title = ?, source_text = ?, updated_at = ? WHERE id = ?",
            (title, text, now(), material_id),
        )
    result = await_adaptation(material_id, text, body.get("language", "kk"))
    return jsonify({"id": material_id, "title": title, "source_text": text, "adaptation": result})


@app.delete("/api/materials/<int:material_id>")
def delete_material(material_id):
    with db() as connection:
        cursor = connection.execute("DELETE FROM materials WHERE id = ?", (material_id,))
    if cursor.rowcount == 0:
        return jsonify({"error": "Материал табылмады."}), 404
    return jsonify({"ok": True})


@app.post("/api/materials/<int:material_id>/progress")
def save_progress(material_id):
    body = request.get_json(silent=True) or {}
    completed = 1 if body.get("completed") else 0
    score = max(0, min(100, int(body.get("score", 0))))
    with db() as connection:
        exists = connection.execute(
            "SELECT id FROM materials WHERE id = ?", (material_id,)
        ).fetchone()
        if not exists:
            return jsonify({"error": "Материал табылмады."}), 404
        connection.execute(
            "INSERT INTO progress(material_id, completed, score, updated_at) VALUES (?, ?, ?, ?) "
            "ON CONFLICT(material_id) DO UPDATE SET completed=excluded.completed, "
            "score=excluded.score, updated_at=excluded.updated_at",
            (material_id, completed, score, now()),
        )
    return jsonify({"completed": completed, "score": score})


@app.post("/api/materials/<int:material_id>/ask")
def ask_about_material(material_id):
    body = request.get_json(silent=True) or {}
    question = clean_text(body.get("question"))
    if not question:
        return jsonify({"error": "Сұрақ жазыңыз."}), 400
    with db() as connection:
        material = connection.execute(
            "SELECT * FROM materials WHERE id = ?", (material_id,)
        ).fetchone()
    if not material:
        return jsonify({"error": "Материал табылмады."}), 404
    answer = local_answer(question, material["source_text"])
    engine = "local"
    if ai_configured():
        try:
            answer = ask_model(
                "Сен Ayla AI, инклюзивті білім беру ассистентісің. Тек берілген материалға сүйен. "
                f"Жауап жоқ болса, оны анық айт. {LANGUAGE_NAMES.get(material['language'], 'қазақ тілі')} тілінде жауап бер.",
                f"Материал:\n{material['source_text']}\n\nСұрақ:\n{question}",
            )
            engine = "ai"
        except Exception:
            pass
    with db() as connection:
        connection.execute(
            "INSERT INTO questions(material_id, question, answer, created_at) VALUES (?, ?, ?, ?)",
            (material_id, question, answer, now()),
        )
    return jsonify({"answer": answer, "engine": engine})


init_db()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "3000")), debug=False)
